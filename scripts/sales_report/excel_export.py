"""Xuất file Excel tổng hợp (mọi người bán): Tổng quan, Theo người, Sản phẩm, Theo ngày, Giao dịch, Ghi chú dữ liệu.

Giá trị ghi sẵn (không dùng công thức) để mở được ở mọi trình xem, kể cả xem trước trong Gmail/điện thoại.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .metrics import WEEKDAYS, prepare

NAVY, BLUE, SOFT, LINE = "1F1B5C", "3949AB", "E8EAF6", "E3E6EE"
STATUS = {"good": ("Đúng/vượt tiến độ", "E1F5EC", "1E7F55"), "warn": ("Hơi chậm", "FCEFD6", "9A6310"),
          "bad": ("Chậm tiến độ", "FBE3E3", "B02A2A"), "none": ("Chưa đặt mục tiêu", "ECEEF3", "6B7089")}
VND, PCT = '#,##0', '0.0%'
thin = Side(style="thin", color=LINE)


def _header(ws, row, titles, widths=None):
    for c, t in enumerate(titles, 1):
        x = ws.cell(row=row, column=c, value=t)
        x.fill, x.font = PatternFill("solid", fgColor=NAVY), Font(bold=True, color="FFFFFF")
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _body(ws, r0, r1, c1, fmts):
    for r in range(r0, r1 + 1):
        for c in range(1, c1 + 1):
            x = ws.cell(row=r, column=c)
            x.border = Border(bottom=thin)
            if fmts.get(c):
                x.number_format = fmts[c]


def _title(ws, text, sub):
    ws["A1"], ws["A2"] = text, sub
    ws["A1"].font, ws["A2"].font = Font(bold=True, size=16, color=NAVY), Font(color="6B7089")
    ws.sheet_view.showGridLines = False


def build_workbook(path: Path, df: pd.DataFrame, m: dict, cfg: dict, issues: list, files_info: list[str]) -> Path:
    rd = dt.date.fromisoformat(m["report_date"])
    gen = cfg["sources"].get("general", {})
    owners = [s["owner"] for s in cfg["sources"]["sources"]]
    d = prepare(df, owners)
    pstart = dt.date.fromisoformat(m["period_start"])
    wstart = min(rd.replace(day=1), pstart)  # tuần vắt qua 2 tháng vẫn đủ dòng
    month = d[(d["date"] >= wstart) & (d["date"] <= rd)].copy()
    wb = Workbook()

    # --- 1. Tổng quan
    ws = wb.active
    ws.title = "Tổng quan"
    _title(ws, f"{gen.get('brand', 'Báo cáo doanh số')} – Tổng hợp doanh số {m['period_title']}",
           f"Tháng {m['month']} · {len(m['people'])} người bán · tạo tự động, số liệu tính bằng Python")
    lw, wtd = m["last_week"], m["week_to_date"]
    kpis = [(f"Doanh thu {m['noun']}", m["day_revenue"], VND), (f"Kỳ trước ({lw['label']})", lw["revenue"], VND),
            ("% so với kỳ trước", (lw["pct"] or 0) / 100 if lw["pct"] is not None else None, PCT),
            ("Doanh thu lũy kế tháng", m["mtd_revenue"], VND), ("Mục tiêu tháng (tổng)", m["total_target"], VND),
            ("% đạt mục tiêu", (m["achieved_pct"] or 0) / 100 if m["achieved_pct"] is not None else None, PCT),
            ("Tiến độ kỳ vọng", m["pace_pct"] / 100, PCT), ("Dự báo cuối tháng", m["projected_month_end"], VND),
            ("Tuần chứa ngày báo cáo", wtd["revenue"], VND), ("Tuần liền trước", wtd["prev_revenue"], VND),
            (f"Số dòng bán trong {m['noun']}", m["day_lines"], '#,##0'), (f"Số đơn vị bán trong {m['noun']}", m["day_units"], '#,##0')]
    _header(ws, 4, ["Chỉ số", "Giá trị"], [38, 22])
    for i, (k, v, f) in enumerate(kpis, 5):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v if v is not None else "n/a")
        c.number_format, c.font, c.alignment = f, Font(bold=True, color=NAVY), Alignment(horizontal="right")
    _body(ws, 5, 4 + len(kpis), 2, {})
    r = 6 + len(kpis)
    ws.cell(row=r, column=1, value=f"Top 3 sản phẩm trong {m['noun']}").font = Font(bold=True, color=NAVY)
    for k, t in enumerate(m["top_products_day"], 1):
        ws.cell(row=r + k, column=1, value=f"{k}. {t['name']}")
        ws.cell(row=r + k, column=2, value=t["revenue"]).number_format = VND

    # --- 2. Theo người
    ws = wb.create_sheet("Theo người")
    _title(ws, "Doanh thu theo người bán", f"{m['period_title']} · tiến độ kỳ vọng {m['pace_pct']:.0f}% · sắp xếp theo doanh thu kỳ")
    heads = ["Hạng", "Người bán", f"Doanh thu {m['noun']}", "Kỳ trước", "% so kỳ trước", "Số dòng", "Đơn vị",
             "Lũy kế tháng", "Mục tiêu tháng", "% đạt mục tiêu", "Trạng thái"]
    _header(ws, 4, heads, [7, 26, 18, 18, 14, 9, 9, 18, 18, 16, 20])
    for i, p in enumerate(m["people"], 5):
        vals = [i - 4, p["name"], p["day_revenue"], p["last_week_day"], None if p["wow_pct"] is None else p["wow_pct"] / 100,
                p["day_lines"], p["day_units"], p["mtd_revenue"], p["target"],
                None if p["achieved_pct"] is None else p["achieved_pct"] / 100, STATUS[p["status"]][0]]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v if v is not None else "n/a")
        _, bg, fg = STATUS[p["status"]]
        st = ws.cell(row=i, column=11)
        st.fill, st.font = PatternFill("solid", fgColor=bg), Font(bold=True, color=fg)
    last = 4 + len(m["people"])
    tot = last + 1
    totals = ["", "TỔNG", m["day_revenue"], lw["revenue"], None if lw["pct"] is None else lw["pct"] / 100, m["day_lines"],
              m["day_units"], m["mtd_revenue"], m["total_target"], None if m["achieved_pct"] is None else m["achieved_pct"] / 100,
              STATUS[m["target_status"]][0]]
    for c, v in enumerate(totals, 1):
        x = ws.cell(row=tot, column=c, value=v if v is not None else "n/a")
        x.font, x.fill = Font(bold=True, color=NAVY), PatternFill("solid", fgColor=SOFT)
    _body(ws, 5, tot, len(heads), {3: VND, 4: VND, 5: PCT, 6: '#,##0', 7: '#,##0', 8: VND, 9: VND, 10: PCT})
    ws.conditional_formatting.add(f"J5:J{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.2, color=BLUE))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(heads))}{last}"

    # --- 3. Sản phẩm
    ws = wb.create_sheet("Sản phẩm")
    _title(ws, "Doanh thu theo sản phẩm", f"{m['period_title']} và lũy kế tháng {m['month']}")
    day = month[(month["date"] >= pstart) & (month["date"] <= rd)]
    g = month.groupby("pkey").agg(name=("product", lambda s: s.mode().iat[0]), rev=("revenue", "sum"), units=("quantity", "sum"))
    gd = day.groupby("pkey").agg(rev=("revenue", "sum"), units=("quantity", "sum")).reindex(g.index).fillna(0)
    g = g.assign(d_rev=gd["rev"], d_units=gd["units"]).sort_values("rev", ascending=False)
    heads = ["Hạng", "Sản phẩm", f"Đơn vị ({m['noun']})", f"Doanh thu {m['noun']}", "Đơn vị (tháng)", "Doanh thu tháng", "Tỷ trọng tháng"]
    _header(ws, 4, heads, [7, 38, 14, 18, 14, 18, 14])
    total_rev = float(g["rev"].sum()) or 1.0
    for i, (_, x) in enumerate(g.iterrows(), 5):
        for c, v in enumerate([i - 4, x["name"], x["d_units"], x["d_rev"], x["units"], x["rev"], x["rev"] / total_rev], 1):
            ws.cell(row=i, column=c, value=v)
    _body(ws, 5, 4 + len(g), len(heads), {3: '#,##0', 4: VND, 5: '#,##0', 6: VND, 7: PCT})
    if len(g):
        ws.conditional_formatting.add(f"G5:G{4 + len(g)}", DataBarRule(start_type="num", start_value=0, end_type="max", color=BLUE))
    ws.freeze_panes = "C5"

    # --- 4. Theo ngày (ma trận ngày × người)
    ws = wb.create_sheet("Theo ngày")
    _title(ws, "Doanh thu từng ngày × từng người", f"Từ {wstart:%d/%m} đến {rd:%d/%m/%Y} · đơn vị: đồng")
    names = [p["name"] for p in m["people"]]
    piv = month.pivot_table(index="date", columns="seller", values="revenue", aggfunc="sum", fill_value=0)
    cols = [n for n in names if n in piv.columns] + [c for c in piv.columns if c not in names]
    heads = ["Ngày", "Thứ"] + cols + ["Tổng ngày"]
    _header(ws, 4, heads, [12, 11] + [15] * len(cols) + [16])
    for i, (dte, row) in enumerate(piv.iterrows(), 5):
        ws.cell(row=i, column=1, value=dte).number_format = "DD/MM/YYYY"
        ws.cell(row=i, column=2, value=WEEKDAYS[dte.weekday()])
        for c, n in enumerate(cols, 3):
            ws.cell(row=i, column=c, value=float(row[n]))
        ws.cell(row=i, column=len(cols) + 3, value=float(row[cols].sum())).font = Font(bold=True, color=NAVY)
    last = 4 + len(piv)
    ws.cell(row=last + 1, column=2, value="TỔNG").font = Font(bold=True, color=NAVY)
    for c in range(3, len(cols) + 4):
        x = ws.cell(row=last + 1, column=c, value=float(piv[cols].sum().sum()) if c == len(cols) + 3 else float(piv[cols[c - 3]].sum()))
        x.font, x.fill = Font(bold=True, color=NAVY), PatternFill("solid", fgColor=SOFT)
    _body(ws, 5, last + 1, len(heads), {c: VND for c in range(3, len(heads) + 1)})
    ws.freeze_panes = "C5"

    # --- 5. Giao dịch
    ws = wb.create_sheet("Giao dịch")
    _title(ws, "Chi tiết giao dịch trong tháng", f"{len(month)} dòng · lọc/sắp xếp bằng nút phễu ở tiêu đề")
    heads = ["Ngày", "Người bán", "Sản phẩm", "Số lượng", "Đơn giá", "Thành tiền", "File nguồn", "Dòng"]
    _header(ws, 4, heads, [12, 24, 36, 10, 14, 16, 28, 8])
    month = month.sort_values(["date", "seller", "row"])
    for i, x in enumerate(month.itertuples(), 5):
        price = x.unit_price if pd.notna(x.unit_price) else (x.revenue / x.quantity if x.quantity else None)
        for c, v in enumerate([x.date, x.seller, x.product, x.quantity, price, x.revenue, x.file, x.row], 1):
            ws.cell(row=i, column=c, value=v)
        ws.cell(row=i, column=1).number_format = "DD/MM/YYYY"
    _body(ws, 5, 4 + len(month), len(heads), {4: '#,##0', 5: VND, 6: VND})
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{4 + len(month)}"

    # --- 6. Ghi chú dữ liệu
    ws = wb.create_sheet("Ghi chú dữ liệu")
    _title(ws, "Kiểm tra dữ liệu đầu vào", "File nguồn đã đọc và các lưu ý chất lượng")
    _header(ws, 4, ["Mục", "Nội dung"], [18, 110])
    row = 5
    for f in files_info:
        ws.cell(row=row, column=1, value="File nguồn"); ws.cell(row=row, column=2, value=f); row += 1
    for i in issues:
        ws.cell(row=row, column=1, value={"critical": "Nghiêm trọng", "warning": "Lưu ý"}.get(i.severity, "Thông tin"))
        ws.cell(row=row, column=2, value=f"{i.title} — {i.detail}"); row += 1
    if not issues:
        ws.cell(row=row, column=1, value="Cảnh báo"); ws.cell(row=row, column=2, value="Không phát hiện bất thường."); row += 1
    for rr in range(5, row):
        ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
